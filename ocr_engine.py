"""Motor de OCR adaptativo para PDFs e imagens.

O módulo não depende de bindings Python do Tesseract: usa o executável oficial
diretamente, mantendo a instalação leve e permitindo que qualquer arquivo
``.traineddata`` instalado no sistema apareça automaticamente na interface.
"""
from __future__ import annotations

import csv
import io
import math
import os
import re
import shutil
import subprocess
import tempfile
import time
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import pymupdf
from PIL import Image, ImageEnhance, ImageFilter, ImageOps, ImageStat

from platform_support import SUBPROCESS_KWARGS


LANGUAGE_NAMES = {
    "afr": "Africâner", "amh": "Amárico", "ara": "Árabe", "aze": "Azerbaijano",
    "bel": "Bielorrusso", "ben": "Bengali", "bod": "Tibetano", "bos": "Bósnio",
    "bre": "Bretão", "bul": "Búlgaro", "cat": "Catalão", "ceb": "Cebuano",
    "ces": "Tcheco", "chi_sim": "Chinês simplificado", "chi_tra": "Chinês tradicional",
    "cym": "Galês", "dan": "Dinamarquês", "deu": "Alemão", "ell": "Grego",
    "eng": "Inglês", "enm": "Inglês médio", "epo": "Esperanto", "est": "Estoniano",
    "eus": "Basco", "fas": "Persa", "fin": "Finlandês", "fra": "Francês",
    "frk": "Fraktur", "frm": "Francês médio", "gle": "Irlandês", "glg": "Galego",
    "grc": "Grego antigo", "guj": "Guzerate", "heb": "Hebraico", "hin": "Hindi",
    "hrv": "Croata", "hun": "Húngaro", "hye": "Armênio", "ind": "Indonésio",
    "isl": "Islandês", "ita": "Italiano", "ita_old": "Italiano antigo", "jpn": "Japonês",
    "kan": "Canarês", "kat": "Georgiano", "kaz": "Cazaque", "khm": "Khmer",
    "kir": "Quirguiz", "kor": "Coreano", "lao": "Laosiano", "lat": "Latim",
    "lav": "Letão", "lit": "Lituano", "mal": "Malaiala", "mar": "Marata",
    "mkd": "Macedônio", "mlt": "Maltês", "msa": "Malaio", "mya": "Birmanês",
    "nep": "Nepalês", "nld": "Holandês", "nor": "Norueguês", "oci": "Occitano",
    "ori": "Oriá", "pan": "Panjabi", "pol": "Polonês", "por": "Português",
    "pus": "Pastó", "ron": "Romeno", "rus": "Russo", "san": "Sânscrito",
    "sin": "Cingalês", "slk": "Eslovaco", "slv": "Esloveno", "spa": "Espanhol",
    "spa_old": "Espanhol antigo", "sqi": "Albanês", "srp": "Sérvio", "swe": "Sueco",
    "syr": "Siríaco", "tam": "Tâmil", "tel": "Telugu", "tgk": "Tadjique",
    "tha": "Tailandês", "tir": "Tigrínia", "tur": "Turco", "uig": "Uigur",
    "ukr": "Ucraniano", "urd": "Urdu", "uzb": "Uzbeque", "vie": "Vietnamita",
    "yid": "Iídiche", "osd": "Detecção de orientação",
}

SUPPORTED_EXTENSIONS = {
    ".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".pnm", ".pgm", ".ppm"
}
MAX_RENDER_PIXELS = 80_000_000


class OCRCancelled(Exception):
    """Interrupção solicitada pelo usuário."""


@dataclass(slots=True)
class OCRSettings:
    engine: str = "adaptive"  # adaptive | tesseract | native
    languages: tuple[str, ...] = ("eng",)
    dpi: int = 300
    layout: str = "auto"  # auto | single | sparse | columns_2 | columns_3 | raw
    preprocessing: str = "auto"  # auto | none | grayscale | contrast | binary
    rotate_pages: bool = True
    deskew: bool = True
    adaptive_threshold: int = 55
    page_expression: str = ""
    preserve_page_breaks: bool = True
    dehyphenate: bool = True
    remove_repeated_margins: bool = True


@dataclass(slots=True)
class PageResult:
    page_number: int
    text: str
    method: str
    confidence: float | None = None
    native_score: int | None = None
    elapsed: float = 0.0


@dataclass(slots=True)
class OCRResult:
    text: str
    pages: list[PageResult] = field(default_factory=list)
    elapsed: float = 0.0

    @property
    def ocr_pages(self) -> int:
        return sum(page.method == "OCR" for page in self.pages)

    @property
    def native_pages(self) -> int:
        return sum(page.method == "Texto nativo" for page in self.pages)

    @property
    def average_confidence(self) -> float | None:
        values = [p.confidence for p in self.pages if p.confidence is not None]
        return sum(values) / len(values) if values else None


def available_tesseract_languages() -> list[str]:
    """Retorna os modelos realmente instalados, sem incluir o modelo OSD."""
    executable = shutil.which("tesseract")
    if not executable:
        return []
    try:
        completed = subprocess.run(
            [executable, "--list-langs"], capture_output=True, text=True, timeout=15, check=False,
            **SUBPROCESS_KWARGS,
        )
        lines = (completed.stdout + "\n" + completed.stderr).splitlines()
        langs = [line.strip() for line in lines if re.fullmatch(r"[A-Za-z0-9_]+", line.strip())]
        return sorted({lang for lang in langs if lang != "osd"})
    except (OSError, subprocess.SubprocessError):
        return []


def language_label(code: str) -> str:
    return f"{LANGUAGE_NAMES.get(code, code)} ({code})"


def parse_page_expression(expression: str, page_count: int) -> list[int]:
    """Converte ``1-3, 7, 10-`` em índices de página (base zero)."""
    value = expression.strip().lower()
    if not value or value in {"todas", "todos", "all", "*"}:
        return list(range(page_count))

    selected: set[int] = set()
    for raw_part in value.split(","):
        part = raw_part.strip()
        if not part:
            continue
        if re.fullmatch(r"\d+", part):
            number = int(part)
            if not 1 <= number <= page_count:
                raise ValueError(f"Página {number} fora do intervalo 1–{page_count}.")
            selected.add(number - 1)
            continue
        match = re.fullmatch(r"(\d*)\s*-\s*(\d*)", part)
        if not match or not any(match.groups()):
            raise ValueError(f"Intervalo de páginas inválido: “{part}”.")
        start = int(match.group(1)) if match.group(1) else 1
        end = int(match.group(2)) if match.group(2) else page_count
        if start < 1 or end > page_count or start > end:
            raise ValueError(f"Intervalo {part} fora de 1–{page_count}.")
        selected.update(range(start - 1, end))
    if not selected:
        raise ValueError("Nenhuma página foi selecionada.")
    return sorted(selected)


def format_page_indices(indices: list[int]) -> str:
    """Formata índices base zero como intervalos base um aceitos pelo OCRmyPDF."""
    if not indices:
        return ""
    ranges: list[str] = []
    start = previous = indices[0] + 1
    for index in indices[1:]:
        current = index + 1
        if current == previous + 1:
            previous = current
            continue
        ranges.append(str(start) if start == previous else f"{start}-{previous}")
        start = previous = current
    ranges.append(str(start) if start == previous else f"{start}-{previous}")
    return ",".join(ranges)


def native_text_quality(text: str) -> int:
    """Pontua de 0 a 100 a utilidade da camada de texto de uma página."""
    text = text.strip()
    if not text:
        return 0
    visible = [char for char in text if not char.isspace()]
    if not visible:
        return 0
    alnum_ratio = sum(char.isalnum() for char in visible) / len(visible)
    replacement_ratio = sum(char in "�□■" for char in visible) / len(visible)
    control_ratio = sum(unicodedata.category(char).startswith("C") for char in visible) / len(visible)
    words = re.findall(r"\b[^\W_]{2,}\b", text, flags=re.UNICODE)
    unique_ratio = len(set(words)) / len(words) if words else 0.0

    length_score = min(1.0, len(visible) / 180.0) * 38
    character_score = min(1.0, alnum_ratio / 0.72) * 32
    word_score = min(1.0, len(words) / 28.0) * 20
    diversity_score = min(1.0, unique_ratio / 0.35) * 10
    penalty = min(55.0, replacement_ratio * 500 + control_ratio * 400)
    return max(0, min(100, round(length_score + character_score + word_score + diversity_score - penalty)))


class OCRProcessor:
    def __init__(
        self,
        settings: OCRSettings,
        progress: Callable[[int, int, str], None] | None = None,
        cancelled: Callable[[], bool] | None = None,
    ):
        self.settings = settings
        self.progress = progress or (lambda _current, _total, _message: None)
        self.cancelled = cancelled or (lambda: False)
        self._active_process: subprocess.Popen | None = None

    def process(self, path: str) -> OCRResult:
        suffix = Path(path).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"Formato não suportado para OCR: {suffix or 'sem extensão'}")
        if not os.path.isfile(path):
            raise FileNotFoundError(path)
        if self.settings.engine != "native" and not shutil.which("tesseract"):
            raise RuntimeError("Tesseract não foi encontrado. Instale o pacote 'tesseract-ocr'.")

        started = time.monotonic()
        pages: list[PageResult] = []
        with pymupdf.open(path) as document:
            indices = parse_page_expression(self.settings.page_expression, document.page_count)
            total = len(indices)
            for position, page_index in enumerate(indices, start=1):
                self._check_cancelled()
                page_number = page_index + 1
                self.progress(position - 1, total, f"Analisando página {page_number}…")
                page = document.load_page(page_index)
                page_started = time.monotonic()
                native = page.get_text("text", sort=True).strip()
                quality = native_text_quality(native)

                use_native = self.settings.engine == "native" or (
                    self.settings.engine == "adaptive" and quality >= self.settings.adaptive_threshold
                )
                if use_native:
                    text = self._extract_native(page)
                    result = PageResult(page_number, text, "Texto nativo", native_score=quality)
                else:
                    text, confidence = self._ocr_page(page)
                    result = PageResult(
                        page_number, text, "OCR", confidence=confidence, native_score=quality
                    )
                result.elapsed = time.monotonic() - page_started
                pages.append(result)
                self.progress(position, total, f"Página {page_number} concluída ({result.method}).")

        page_texts = [self._clean_page(page.text) for page in pages]
        if self.settings.remove_repeated_margins and len(page_texts) >= 3:
            page_texts = self._remove_repeated_margins(page_texts)
        for page, text in zip(pages, page_texts):
            page.text = text
        joined = self._join_pages(pages)
        return OCRResult(joined.strip(), pages, time.monotonic() - started)

    def cancel(self) -> None:
        process = self._active_process
        if process is not None and process.poll() is None:
            process.terminate()

    def _check_cancelled(self) -> None:
        if self.cancelled():
            self.cancel()
            raise OCRCancelled("OCR cancelado pelo usuário.")

    def _extract_native(self, page: pymupdf.Page) -> str:
        blocks = [block for block in page.get_text("blocks") if len(block) >= 5 and block[4].strip()]
        if not blocks:
            return page.get_text("text", sort=True).strip()
        columns = {"columns_2": 2, "columns_3": 3}.get(self.settings.layout)
        if not columns:
            return page.get_text("text", sort=True).strip()

        width = max(float(page.rect.width), 1.0)
        ordered: list[tuple] = []
        for column in range(columns):
            left, right = width * column / columns, width * (column + 1) / columns
            candidates = [b for b in blocks if left <= (b[0] + b[2]) / 2 < right]
            ordered.extend(sorted(candidates, key=lambda b: (round(b[1] / 4), b[0])))
        return "\n\n".join(str(block[4]).strip() for block in ordered)

    def _ocr_page(self, page: pymupdf.Page) -> tuple[str, float | None]:
        scale = self.settings.dpi / 72.0
        projected_pixels = page.rect.width * scale * page.rect.height * scale
        if projected_pixels > MAX_RENDER_PIXELS:
            scale *= math.sqrt(MAX_RENDER_PIXELS / projected_pixels)
        pixmap = page.get_pixmap(
            matrix=pymupdf.Matrix(scale, scale), alpha=False, colorspace=pymupdf.csRGB
        )
        image = Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB")

        if self.settings.rotate_pages:
            image = self._orient_image(image)
        image = self._preprocess_image(image)

        columns = {"columns_2": 2, "columns_3": 3}.get(self.settings.layout, 1)
        if columns > 1:
            texts: list[str] = []
            confidences: list[float] = []
            overlap = max(8, image.width // 250)
            for column in range(columns):
                self._check_cancelled()
                left = max(0, round(image.width * column / columns) - (overlap if column else 0))
                right = min(
                    image.width,
                    round(image.width * (column + 1) / columns) + (overlap if column + 1 < columns else 0),
                )
                text, confidence = self._run_tesseract(image.crop((left, 0, right, image.height)), 6)
                texts.append(text.strip())
                if confidence is not None:
                    confidences.append(confidence)
            return "\n\n".join(filter(None, texts)), (sum(confidences) / len(confidences) if confidences else None)

        psm = {"auto": 3, "single": 6, "sparse": 11, "raw": 13}.get(self.settings.layout, 3)
        return self._run_tesseract(image, psm)

    def _orient_image(self, image: Image.Image) -> Image.Image:
        if "osd" not in self._all_tesseract_languages():
            return image
        # OSD não precisa da imagem em resolução total e fica muito mais rápido assim.
        probe = image.copy()
        probe.thumbnail((1800, 1800))
        try:
            output = self._run_command(probe, ["-l", "osd", "--psm", "0"])
            match = re.search(r"Rotate:\s*(90|180|270)", output)
            if match:
                return image.rotate(-int(match.group(1)), expand=True, fillcolor="white")
        except (RuntimeError, OCRCancelled):
            if self.cancelled():
                raise
        return image

    def _all_tesseract_languages(self) -> set[str]:
        executable = shutil.which("tesseract")
        if not executable:
            return set()
        try:
            result = subprocess.run(
                [executable, "--list-langs"], capture_output=True, text=True, timeout=10, **SUBPROCESS_KWARGS
            )
            return set(re.findall(r"(?m)^[A-Za-z0-9_]+$", result.stdout))
        except (OSError, subprocess.SubprocessError):
            return set()

    def _preprocess_image(self, image: Image.Image) -> Image.Image:
        mode = self.settings.preprocessing
        if mode == "none":
            return image
        gray = ImageOps.grayscale(image)
        if mode == "grayscale":
            return self._deskew(gray) if self.settings.deskew else gray

        gray = ImageOps.autocontrast(gray, cutoff=1)
        if mode == "contrast":
            gray = ImageEnhance.Contrast(gray).enhance(1.45)
        elif mode == "binary":
            threshold = self._otsu_threshold(gray)
            gray = gray.point(lambda pixel: 255 if pixel > threshold else 0, mode="1")
        else:  # auto
            stat = ImageStat.Stat(gray)
            deviation = stat.stddev[0] if stat.stddev else 0
            if deviation < 45:
                gray = ImageEnhance.Contrast(gray).enhance(1.35)
            gray = gray.filter(ImageFilter.SHARPEN)
        return self._deskew(gray) if self.settings.deskew else gray

    @staticmethod
    def _otsu_threshold(image: Image.Image) -> int:
        histogram = image.histogram()[:256]
        total = sum(histogram)
        weighted_sum = sum(i * count for i, count in enumerate(histogram))
        background_weight = 0
        background_sum = 0
        best_variance = -1.0
        threshold = 127
        for value, count in enumerate(histogram):
            background_weight += count
            if not background_weight:
                continue
            foreground_weight = total - background_weight
            if not foreground_weight:
                break
            background_sum += value * count
            mean_background = background_sum / background_weight
            mean_foreground = (weighted_sum - background_sum) / foreground_weight
            variance = background_weight * foreground_weight * (mean_background - mean_foreground) ** 2
            if variance > best_variance:
                best_variance = variance
                threshold = value
        return threshold

    def _deskew(self, image: Image.Image) -> Image.Image:
        """Corrige inclinações pequenas pela variância da projeção horizontal."""
        probe = image.convert("L")
        probe.thumbnail((1100, 1100))
        # Evita gastar CPU em páginas quase vazias.
        if ImageStat.Stat(probe).stddev[0] < 12:
            return image
        best_angle = 0.0
        best_score = -1.0
        for angle in (-3.0, -2.0, -1.0, 0.0, 1.0, 2.0, 3.0):
            rotated = probe.rotate(angle, expand=False, fillcolor=255)
            width, height = rotated.size
            pixels = rotated.load()
            rows = [sum(255 - pixels[x, y] for x in range(0, width, 3)) for y in range(0, height, 2)]
            mean = sum(rows) / max(len(rows), 1)
            score = sum((value - mean) ** 2 for value in rows) / max(len(rows), 1)
            if score > best_score:
                best_score, best_angle = score, angle
        if abs(best_angle) < 0.1:
            return image
        return image.rotate(best_angle, expand=True, fillcolor=255)

    def _run_tesseract(self, image: Image.Image, psm: int) -> tuple[str, float | None]:
        languages = "+".join(self.settings.languages) or "eng"
        tsv = self._run_command(
            image,
            ["-l", languages, "--psm", str(psm), "-c", "preserve_interword_spaces=1", "tsv"],
        )
        return self._parse_tsv(tsv)

    def _run_command(self, image: Image.Image, arguments: list[str]) -> str:
        executable = shutil.which("tesseract")
        if not executable:
            raise RuntimeError("Tesseract não foi encontrado no sistema.")
        buffer = io.BytesIO()
        image.save(buffer, format="PNG", optimize=False)
        command = [executable, "stdin", "stdout", *arguments]
        self._active_process = subprocess.Popen(
            command, stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            **SUBPROCESS_KWARGS,
        )
        assert self._active_process.stdin is not None
        try:
            stdout, stderr = self._active_process.communicate(buffer.getvalue(), timeout=300)
        except subprocess.TimeoutExpired:
            self._active_process.kill()
            self._active_process.communicate()
            raise RuntimeError("O Tesseract excedeu o limite de 5 minutos nesta página.")
        finally:
            process = self._active_process
            self._active_process = None
        self._check_cancelled()
        if process.returncode != 0:
            details = stderr.decode("utf-8", errors="replace").strip()
            raise RuntimeError(f"Tesseract falhou: {details or 'erro desconhecido'}")
        return stdout.decode("utf-8", errors="replace")

    @staticmethod
    def _parse_tsv(tsv: str) -> tuple[str, float | None]:
        reader = csv.DictReader(io.StringIO(tsv), delimiter="\t")
        lines: list[str] = []
        confidences: list[float] = []
        current_key: tuple[str, str, str, str] | None = None
        current_words: list[str] = []
        previous_block: str | None = None
        for row in reader:
            word = (row.get("text") or "").strip()
            if not word:
                continue
            key = (
                row.get("page_num", ""), row.get("block_num", ""),
                row.get("par_num", ""), row.get("line_num", ""),
            )
            if current_key is not None and key != current_key:
                if previous_block is not None and key[1] != previous_block and lines:
                    lines.append("")
                lines.append(" ".join(current_words))
                previous_block = current_key[1]
                current_words = []
            current_key = key
            current_words.append(word)
            try:
                confidence = float(row.get("conf", "-1"))
                if confidence >= 0:
                    confidences.append(confidence)
            except ValueError:
                pass
        if current_words:
            lines.append(" ".join(current_words))
        text = "\n".join(lines)
        return text, (sum(confidences) / len(confidences) if confidences else None)

    def _clean_page(self, text: str) -> str:
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        text = re.sub(r"[ \t]+\n", "\n", text)
        text = re.sub(r"\n{4,}", "\n\n\n", text)
        if self.settings.dehyphenate:
            text = re.sub(r"(?<=\w)-\n(?=[a-zà-öø-ÿ])", "", text, flags=re.IGNORECASE)
        return text.strip()

    @staticmethod
    def _line_signature(line: str) -> str:
        normalized = re.sub(r"\d+", "#", line.casefold())
        return re.sub(r"\W+", " ", normalized).strip()

    def _remove_repeated_margins(self, pages: list[str]) -> list[str]:
        signatures: Counter[str] = Counter()
        per_page: list[list[str]] = []
        for text in pages:
            lines = text.splitlines()
            nonempty_indices = [i for i, line in enumerate(lines) if line.strip()]
            margin_indices = set(nonempty_indices[:2] + nonempty_indices[-2:])
            page_signatures = {
                self._line_signature(lines[i]) for i in margin_indices if len(lines[i].strip()) <= 160
            }
            signatures.update(signature for signature in page_signatures if signature)
            per_page.append(lines)
        minimum = max(2, math.ceil(len(pages) * 0.5))
        repeated = {signature for signature, count in signatures.items() if count >= minimum}
        cleaned = []
        for lines in per_page:
            nonempty_indices = [i for i, line in enumerate(lines) if line.strip()]
            margin_indices = set(nonempty_indices[:2] + nonempty_indices[-2:])
            kept = [
                line for index, line in enumerate(lines)
                if index not in margin_indices or self._line_signature(line) not in repeated
            ]
            cleaned.append("\n".join(kept).strip())
        return cleaned

    def _join_pages(self, pages: list[PageResult]) -> str:
        if self.settings.preserve_page_breaks:
            return "\n\n".join(f"--- Página {page.page_number} ---\n\n{page.text}" for page in pages)
        return "\n\n".join(page.text for page in pages)


def export_docx(path: str, result: OCRResult) -> None:
    """Salva o resultado preservando páginas e parágrafos básicos."""
    import docx

    document = docx.Document()
    for index, page in enumerate(result.pages):
        if index:
            document.add_page_break()
        document.add_heading(f"Página {page.page_number}", level=2)
        for paragraph in re.split(r"\n\s*\n", page.text):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
    document.save(path)


def create_searchable_pdf(
    source: str,
    destination: str,
    languages: tuple[str, ...],
    dpi: int = 300,
    force_ocr: bool = False,
    cancelled: Callable[[], bool] | None = None,
    page_expression: str = "",
) -> None:
    """Cria PDF pesquisável via OCRmyPDF, quando disponível."""
    executable = shutil.which("ocrmypdf")
    if not executable:
        raise RuntimeError("OCRmyPDF não foi encontrado no sistema.")
    if Path(source).suffix.lower() != ".pdf":
        raise ValueError("A criação de PDF pesquisável requer um PDF como origem.")
    command = [
        executable, "--quiet", "--output-type", "pdf", "--optimize", "1", "--deskew",
        "--rotate-pages", "--oversample", str(dpi), "--language", "+".join(languages),
    ]
    if page_expression.strip() and page_expression.strip().lower() not in {"todas", "todos", "all", "*"}:
        with pymupdf.open(source) as document:
            normalized_pages = format_page_indices(
                parse_page_expression(page_expression, document.page_count)
            )
        command.extend(["--pages", normalized_pages])
    command.append("--force-ocr" if force_ocr else "--skip-text")
    destination_path = Path(destination)
    descriptor, temporary = tempfile.mkstemp(
        prefix=f".{destination_path.stem}_", suffix=".pdf", dir=str(destination_path.parent)
    )
    os.close(descriptor)
    os.remove(temporary)  # OCRmyPDF exige que a saída ainda não exista.
    command.extend([source, temporary])
    process = subprocess.Popen(
        command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, **SUBPROCESS_KWARGS
    )
    started = time.monotonic()
    while True:
        if cancelled is not None and cancelled():
            process.terminate()
            try:
                process.communicate(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill()
                process.communicate()
            if os.path.exists(temporary):
                os.remove(temporary)
            raise OCRCancelled("Criação do PDF pesquisável cancelada.")
        try:
            stdout, stderr = process.communicate(timeout=1)
            break
        except subprocess.TimeoutExpired:
            if time.monotonic() - started > 3600:
                process.kill()
                process.communicate()
                if os.path.exists(temporary):
                    os.remove(temporary)
                raise RuntimeError("OCRmyPDF excedeu o limite de uma hora.")
    if process.returncode != 0:
        details = (stderr or stdout).strip()
        if os.path.exists(temporary):
            os.remove(temporary)
        raise RuntimeError(f"OCRmyPDF falhou: {details}")
    os.replace(temporary, destination)
